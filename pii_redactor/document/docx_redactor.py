from docx import Document
from detectors.pii_detector import PIIDetector
from replacers.replacement_manager import ReplacementManager

def process_text_elements(elements, detector, rep_manager):
    for element in elements:
        text = element.text
        if not text.strip():
            continue
            
        entities = detector.detect(text)
        if not entities:
            continue
            
        # Replace right-to-left
        entities.sort(key=lambda x: x['start'], reverse=True)
        new_text = text
        for ent in entities:
            fake_val = rep_manager.get_replacement(ent['value'], ent['type'])
            new_text = new_text[:ent['start']] + fake_val + new_text[ent['end']:]
            
        element.text = new_text

def redact_document(input_path, output_path):
    doc = Document(input_path)
    detector = PIIDetector()
    rep_manager = ReplacementManager()
    
    # Process paragraphs
    for p in doc.paragraphs:
        process_text_elements([p], detector, rep_manager)
        
    # Process tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    process_text_elements([p], detector, rep_manager)
                    
    # Save
    doc.save(output_path)
    rep_manager.save_mapping()
